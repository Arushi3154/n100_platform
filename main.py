import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill
from src.screener.engine import ScreenerEngine
from src.analytics.peer import PeerAnalyticsEngine
from src.analytics.radar import generate_radar_charts

def generate_synthetic_universe():
    np.random.seed(42)
    companies = [f"CMP_{i:02d}" for i in range(1, 93)]
    sectors = ['IT Services', 'Financials', 'FMCG', 'Automobile', 'Pharma', 'Energy']
    peers = [f"Peer_Group_{i}" for i in range(1, 12)]

    data = []
    peer_mappings = []

    for idx, c in enumerate(companies):
        sector = np.random.choice(sectors)
        peer = np.random.choice(peers) if idx < 85 else None
        
        data.append({
            'company_id': c,
            'company_name': f"Company {c}",
            'broad_sector': sector,
            'roe': np.random.uniform(5, 30),
            'roce': np.random.uniform(5, 35),
            'npm': np.random.uniform(3, 25),
            'opm': np.random.uniform(5, 30),
            'de': 0.0 if idx % 10 == 0 else np.random.uniform(0.1, 2.5),
            'fcf': np.random.uniform(-300, 1500),
            'fcf_cagr': np.random.uniform(-5, 25),
            'cfo_pat_ratio': np.random.uniform(0.5, 1.5),
            'revenue_cagr_5yr': np.random.uniform(2, 25),
            'revenue_cagr_3yr': np.random.uniform(2, 25),
            'pat_cagr_5yr': np.random.uniform(2, 30),
            'pe': np.random.uniform(8, 42),
            'pb': np.random.uniform(0.8, 5.5),
            'dividend_yield': np.random.uniform(0, 4.5),
            'dividend_payout': np.random.uniform(10, 85),
            'icr': 'Debt Free' if idx % 10 == 0 else str(np.random.uniform(1.5, 20)),
            'sales': np.random.uniform(1000, 20000),
            'market_cap': np.random.uniform(2000, 100000),
            'net_profit': np.random.uniform(100, 5000),
            'asset_turnover': np.random.uniform(0.5, 2.5),
            'eps_cagr_5yr': np.random.uniform(3, 25)
        })
        if peer:
            peer_mappings.append({'company_id': c, 'peer_group_name': peer})

    return pd.DataFrame(data), pd.DataFrame(peer_mappings)

def export_styled_screener(df, presets, engine):
    green_fill = PatternFill(start_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", fill_type="solid")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for preset_name in presets:
        res = engine.run_preset(df, preset_name)
        ws = wb.create_sheet(title=preset_name)
        
        headers = list(res.columns)
        ws.append(headers)

        preset_rules = engine.config['presets'].get(preset_name, {})

        for _, row in res.iterrows():
            ws.append(list(row.values))
            curr_row = ws.max_row
            
            for col_idx, col_name in enumerate(headers, 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                if col_name == 'roe' and 'roe_min' in preset_rules:
                    cell.fill = green_fill if cell.value >= preset_rules['roe_min'] else red_fill
                elif col_name == 'de' and 'de_max' in preset_rules:
                    cell.fill = green_fill if cell.value <= preset_rules['de_max'] else red_fill
                elif col_name == 'fcf' and 'fcf_min' in preset_rules:
                    cell.fill = green_fill if cell.value >= preset_rules['fcf_min'] else red_fill

    wb.save("output/screener_output.xlsx")

def build_peer_excel(df, peer_groups_df, pcts_df):
    merged = df.merge(peer_groups_df, on="company_id", how="left")
    groups = peer_groups_df['peer_group_name'].unique()
    
    amber_fill = PatternFill(start_color="FFBF00", fill_type="solid")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for grp in groups:
        grp_data = merged[merged['peer_group_name'] == grp].copy()
        if grp_data.empty:
            continue

        ws = wb.create_sheet(title=grp)
        ws.append(list(grp_data.columns))

        for row_idx, row in grp_data.reset_index(drop=True).iterrows():
            ws.append(list(row.values))
            curr_row = ws.max_row
            
            if row_idx == 0:
                for cell in ws[curr_row]:
                    cell.fill = amber_fill

        median_vals = ["MEDIAN", "Group Median"] + [
            round(float(grp_data[col].median()), 2) if pd.api.types.is_numeric_dtype(grp_data[col]) else "" 
            for col in grp_data.columns[2:]
        ]
        ws.append(median_vals)

    wb.save("output/peer_comparison.xlsx")

if __name__ == "__main__":
    df, peer_groups_df = generate_synthetic_universe()
    screener = ScreenerEngine()
    peer_engine = PeerAnalyticsEngine()

    presets = list(screener.config['presets'].keys())
    export_styled_screener(df, presets, screener)
    print("✓ Screener workbook generated: output/screener_output.xlsx")

    df_with_icr = screener.apply_icr_and_de_rules(df)
    pcts_df = peer_engine.compute_percentiles(df_with_icr, peer_groups_df)
    print("✓ Peer percentiles stored in SQLite database: data/n100_platform.db")

    full_df = screener.compute_composite_score(df_with_icr)
    generate_radar_charts(full_df, peer_groups_df)
    print("✓ Radar charts exported to: reports/radar_charts/")

    build_peer_excel(full_df, peer_groups_df, pcts_df)
    print("✓ Peer comparison report generated: output/peer_comparison.xlsx")
