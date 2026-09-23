import yaml
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill
from typing import Dict, Any, Union, Tuple

FINANCIAL_SECTORS = ["FINANCIALS", "FINANCIAL SERVICES", "BANKS", "NBFC"]

def apply_winsorisation(s: pd.Series, lower_p: float = 10.0, upper_p: float = 90.0) -> pd.Series:
    clean_s = s.dropna()
    if len(clean_s) == 0:
        return s
    p_low = np.percentile(clean_s, lower_p)
    p_high = np.percentile(clean_s, upper_p)
    return s.clip(lower=p_low, upper=p_high)

class ScreenerEngine:
    def __init__(self, config_path: str = "config/screener_config.yaml"):
        self.config_path = config_path
        self.config = self.load_config()

    def load_config(self) -> Dict[str, Any]:
        try:
            with open(self.config_path, "r") as f:
                return yaml.safe_load(f)
        except Exception:
            return {"presets": {}, "composite_weights": {}}

    def calculate_composite_score(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if df.empty:
            df["composite_quality_score"] = []
            return df

        for col in ["return_on_equity_pct", "revenue_cagr_5yr", "free_cash_flow_cr"]:
            if col in df.columns:
                cleaned = df[col].fillna(0)
                winsorized = apply_winsorisation(cleaned)
                min_v, max_v = winsorized.min(), winsorized.max()
                if max_v == min_v:
                    df[f"{col}_score"] = 50.0
                else:
                    df[f"{col}_score"] = ((winsorized - min_v) / (max_v - min_v)) * 100.0
            else:
                df[f"{col}_score"] = 50.0

        df["composite_quality_score"] = (
            df.get("return_on_equity_pct_score", 50.0) * 0.35 +
            df.get("free_cash_flow_cr_score", 50.0) * 0.30 +
            df.get("revenue_cagr_5yr_score", 50.0) * 0.20 +
            50.0 * 0.15
        ).round(2)
        return df

    def run_preset(self, df: pd.DataFrame, preset_rules_or_name: Union[Dict[str, Any], str]) -> Union[pd.DataFrame, Tuple[pd.DataFrame, pd.DataFrame]]:
        if isinstance(preset_rules_or_name, str):
            preset_rules = self.config.get("presets", {}).get(preset_rules_or_name, {})
        elif isinstance(preset_rules_or_name, dict):
            preset_rules = preset_rules_or_name
        else:
            preset_rules = {}

        filtered = df.copy()
        
        for metric, bounds in preset_rules.items():
            if isinstance(bounds, (list, tuple)) and len(bounds) == 2:
                min_val, max_val = bounds
                if metric in filtered.columns:
                    if min_val is not None:
                        if metric == "interest_coverage":
                            mask = (filtered[metric] >= min_val) | (filtered[metric].isna())
                            filtered = filtered[mask]
                        else:
                            filtered = filtered[filtered[metric] >= min_val]
                    if max_val is not None:
                        if metric == "debt_to_equity":
                            is_fin = filtered["broad_sector"].astype(str).str.upper().isin(FINANCIAL_SECTORS)
                            mask = (filtered[metric] <= max_val) | is_fin
                            filtered = filtered[mask]
                        else:
                            filtered = filtered[filtered[metric] <= max_val]
            elif metric.endswith("_min"):
                col = metric[:-4]
                if col in filtered.columns:
                    if col == "interest_coverage":
                        mask = (filtered[col] >= bounds) | (filtered[col].isna())
                        filtered = filtered[mask]
                    else:
                        filtered = filtered[filtered[col] >= bounds]
            elif metric.endswith("_max"):
                col = metric[:-4]
                if col in filtered.columns:
                    if col == "debt_to_equity":
                        is_fin = filtered["broad_sector"].astype(str).str.upper().isin(FINANCIAL_SECTORS)
                        mask = (filtered[col] <= bounds) | is_fin
                        filtered = filtered[mask]
                    else:
                        filtered = filtered[filtered[col] <= bounds]

        filtered = self.calculate_composite_score(filtered)
        filtered = filtered.sort_values(by="composite_quality_score", ascending=False)
        return filtered

    def export_excel(self, preset_results: Dict[str, pd.DataFrame], output_path: str = "output/screener_output.xlsx"):
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for preset_name, df in preset_results.items():
                sheet_title = preset_name[:30]
                df.to_excel(writer, sheet_name=sheet_title, index=False)
                ws = writer.sheets[sheet_title]
                green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=ws.max_column).fill = green_fill

def load_screener_config(config_path: str = "config/screener_config.yaml") -> Dict[str, Any]:
    engine = ScreenerEngine(config_path=config_path)
    return engine.config

def run_screener_preset(df: pd.DataFrame, preset_rules: Dict[str, Any]) -> pd.DataFrame:
    engine = ScreenerEngine()
    return engine.run_preset(df, preset_rules)

def calculate_composite_score(df: pd.DataFrame) -> pd.DataFrame:
    engine = ScreenerEngine()
    return engine.calculate_composite_score(df)

def export_screener_excel(preset_results: Dict[str, pd.DataFrame], output_path: str = "output/screener_output.xlsx"):
    engine = ScreenerEngine()
    engine.export_excel(preset_results, output_path)
